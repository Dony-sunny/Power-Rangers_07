import csv
import re
import zipfile
from io import BytesIO, StringIO
from datetime import datetime
from fastapi import HTTPException
from openpyxl import load_workbook, Workbook
from pydantic import ValidationError
from backend.schemas.requests import CargoCreate

HEADERS = [
    "reference",
    "cargo_type",
    "weight_tonnes",
    "volume_m3",
    "packaging",
    "origin",
    "destination",
    "ready_time",
    "deadline",
    "fragile",
    "hazardous",
    "consolidation_allowed",
]
ALIASES = {
    "cargo": "cargo_type",
    "commodity": "cargo_type",
    "weight": "weight_tonnes",
    "weight_t": "weight_tonnes",
    "tonnes": "weight_tonnes",
    "volume": "volume_m3",
    "from": "origin",
    "to": "destination",
    "deadline": "delivery_deadline",
    "required_date": "delivery_deadline",
    "ready": "ready_time",
    "ref": "reference",
}


def header(value):
    normalized = re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")
    return ALIASES.get(normalized, normalized)


def read_rows(content, suffix):
    if suffix == ".csv":
        from intelligence.document_intake.documents import csv_preview

        return csv_preview(content)
    if suffix != ".xlsx":
        raise HTTPException(415, "Bulk intake supports CSV or XLSX.")
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if (
                len(archive.infolist()) > 1000
                or sum(item.file_size for item in archive.infolist()) > 32 * 1024 * 1024
            ):
                raise HTTPException(413, "Workbook decompressed size exceeds 32 MB.")
            if any(
                "vbaProject" in item.filename or "externalLinks/" in item.filename
                for item in archive.infolist()
            ):
                raise HTTPException(
                    415, "Macros and external workbook links are not accepted."
                )
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            sheet = workbook.active
            if (
                sheet.max_row
                and sheet.max_row > 201
                or sheet.max_column
                and sheet.max_column > 40
            ):
                raise HTTPException(413, "Workbook limit: 200 rows and 40 columns.")
            iterator = sheet.iter_rows(max_row=202, max_col=40, values_only=True)
            headers = list(next(iterator))
            mapped = [header(key) for key in headers if key is not None]
            if len(set(mapped)) != len(mapped):
                raise HTTPException(
                    422,
                    "Duplicate mapped headers are ambiguous. Use one column per field.",
                )
            return [
                {
                    str(key): value
                    for key, value in zip(headers, values)
                    if key is not None
                }
                for values in iterator
                if any(value is not None for value in values)
            ]
        finally:
            workbook.close()
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            422, "Workbook could not be safely read. Use the supplied XLSX template."
        ) from None


def preview_rows(content, suffix):
    results = []
    for number, raw in enumerate(read_rows(content, suffix), start=2):
        fields, warnings, reference, unsafe = {}, [], "", []
        for original, value in raw.items():
            key = header(original)
            if isinstance(value, str) and value.startswith("="):
                unsafe.append(
                    {
                        "loc": [key],
                        "msg": "Formula cells are not accepted; supply explicit values.",
                        "type": "value_error",
                    }
                )
                continue
            if value is None or value == "":
                continue
            if key == "reference":
                reference = str(value)[:100]
                continue
            if key not in CargoCreate.model_fields:
                warnings.append(f"Ignored unrecognized column: {original}")
                continue
            if isinstance(value, datetime):
                value = value.isoformat()
            if key in fields:
                warnings.append(f"Duplicate mapped column: {key}")
            fields[key] = value
        try:
            if unsafe:
                results.append(
                    {
                        "row": number,
                        "reference": reference,
                        "valid": False,
                        "errors": unsafe,
                        "warnings": warnings,
                    }
                )
                continue
            parsed = CargoCreate.model_validate(fields)
            results.append(
                {
                    "row": number,
                    "reference": reference,
                    "valid": True,
                    "fields": parsed.model_dump(mode="json"),
                    "warnings": warnings,
                }
            )
        except ValidationError as error:
            results.append(
                {
                    "row": number,
                    "reference": reference,
                    "valid": False,
                    "errors": error.errors(
                        include_url=False, include_context=False, include_input=False
                    ),
                    "warnings": warnings,
                }
            )
    return results


def template_bytes():
    from data.seed.network import tomorrow
    from datetime import timedelta

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cargo"
    sheet.append(HEADERS)
    start = tomorrow()
    sheet.append(
        [
            "PO-001",
            "cement",
            80,
            60,
            "bagged",
            "Kalamassery",
            "Alappuzha",
            start.isoformat(),
            (start + timedelta(hours=14)).isoformat(),
            False,
            False,
            True,
        ]
    )
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 24
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def error_csv(rows):
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(["row", "reference", "valid", "errors", "warnings"])
    for row in rows:
        messages = "; ".join(
            ".".join(map(str, error["loc"])) + ": " + error["msg"]
            for error in row.get("errors", [])
        )
        values = [
            row["row"],
            row.get("reference", ""),
            row["valid"],
            messages,
            "; ".join(row.get("warnings", [])),
        ]
        writer.writerow(
            [
                "'" + str(value)
                if str(value).startswith(("=", "+", "-", "@"))
                else value
                for value in values
            ]
        )
    return output.getvalue()
