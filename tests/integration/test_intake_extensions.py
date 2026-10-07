from io import BytesIO
from openpyxl import load_workbook
from PIL import Image
from intelligence.document_intake.spreadsheets import template_bytes


def test_xlsx_review_import_valid_only_and_reuse_guard(client):
    workbook = load_workbook(BytesIO(template_bytes()))
    sheet = workbook.active
    row = [cell.value for cell in sheet[2]]
    row[2] = -10
    sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    response = client.post(
        "/api/intake/cargo/bulk-preview",
        files={
            "file": (
                "test.xlsx",
                output.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200, response.text
    preview = response.json()
    assert (preview["valid_count"], preview["invalid_count"]) == (1, 1)
    endpoint = f"/api/intake/cargo/bulk-import/{preview['preview_id']}"
    assert client.post(endpoint).status_code == 422
    assert client.post(endpoint + "?approved=true").json()["imported_count"] == 1
    assert client.post(endpoint + "?approved=true").status_code == 409
    assert (
        client.get(
            f"/api/intake/cargo/bulk-preview/{preview['preview_id']}/errors"
        ).status_code
        == 200
    )


def test_xlsx_formulas_do_not_execute(client):
    workbook = load_workbook(BytesIO(template_bytes()))
    workbook.active["C2"] = "=1+79"
    output = BytesIO()
    workbook.save(output)
    response = client.post(
        "/api/intake/cargo/bulk-preview",
        files={"file": ("formula.xlsx", output.getvalue())},
    )
    assert response.json()["invalid_count"] == 1


def test_image_ocr_review_and_fallback(client, monkeypatch):
    output = BytesIO()
    Image.new("RGB", (100, 100), "white").save(output, "PNG")
    monkeypatch.setattr(
        "backend.api.intake.run_ocr",
        lambda *args: {
            "text": "80 tonnes cement bagged from Kochi to Alappuzha",
            "source": "mock_ocr",
            "warnings": [],
            "confidence": 0.9,
        },
    )
    response = client.post(
        "/api/intake/cargo/document",
        files={"file": ("../scan.png", output.getvalue(), "image/png")},
    )
    assert response.status_code == 200
    extraction = response.json()
    assert extraction["fields"]["weight_tonnes"] == 80
    assert "ready_time" in extraction["missing_fields"]
    assert extraction["requires_confirmation"]
    assert extraction["filename"] == "scan.png"
    assert (
        client.post(
            "/api/intake/cargo/document", files={"file": ("fake.png", b"not an image")}
        ).status_code
        == 415
    )


def test_ocr_timeout_returns_empty_fields(monkeypatch):
    import subprocess
    from intelligence.document_intake.ocr import run_ocr

    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired("ocr", 20)

    monkeypatch.setattr(subprocess, "run", timeout)
    assert run_ocr(b"fake", ".png")["source"] == "ocr_unavailable"
